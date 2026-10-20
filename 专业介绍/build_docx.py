# -*- coding: utf-8 -*-
"""生成《广东工业大学机器人专业介绍》Word 文档。

数据来源：
- 学校/学院官方页面（逐段标注原链接）
- images/ 目录下已下载图片（每张标注原图链接与来源页面）
- 论文佐证.md（28 篇论文，DOI 链接）
"""
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from PIL import Image

BASE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(BASE, "images")
OUT = os.path.join(BASE, "广东工业大学机器人专业介绍.docx")

CN_FONT = "微软雅黑"
CN_BODY = "宋体"
EN_FONT = "Times New Roman"
GRAY = RGBColor(0x59, 0x59, 0x59)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
LINK = RGBColor(0x05, 0x63, 0xC1)


# ---------------------------------------------------------------- 基础工具
def set_run_font(run, cn=CN_BODY, en=EN_FONT, size=None, bold=None, color=None):
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)
    rfonts.set(qn("w:eastAsia"), cn)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_hyperlink(paragraph, url, text, size=9):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), EN_FONT)
    rfonts.set(qn("w:hAnsi"), EN_FONT)
    rfonts.set(qn("w:eastAsia"), CN_BODY)
    rpr.append(rfonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    new_run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hl.append(new_run)
    paragraph._p.append(hl)
    return hl


def para(doc, text="", size=10.5, bold=False, align=None, indent=True,
         space_after=6, cn=CN_BODY, color=None, line=1.45):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing = line
    if indent:
        pf.first_line_indent = Pt(size * 2)
    if text:
        r = p.add_run(text)
        set_run_font(r, cn=cn, size=size, bold=bold, color=color)
    return p


def source_line(doc, label, url):
    """一段小字来源标注：来源：xxx  <超链接>"""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(label)
    set_run_font(r, size=8.5, color=GRAY)
    if url:
        add_hyperlink(p, url, url, size=8.5)
    return p


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, cn=CN_FONT, size=17, bold=True, color=BLUE)
    # 下边框
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "1F4E79")
    pbdr.append(bottom)
    ppr.append(pbdr)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, cn=CN_FONT, size=13, bold=True, color=RGBColor(0x2E, 0x5C, 0x8A))
    return p


def h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, cn=CN_FONT, size=11, bold=True, color=RGBColor(0x33, 0x33, 0x33))
    return p


def bullet(doc, text, size=10.5, bold_head=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Pt(18)
    pf.first_line_indent = Pt(-12)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.4
    r = p.add_run("▪ ")
    set_run_font(r, size=size)
    if bold_head:
        r2 = p.add_run(bold_head)
        set_run_font(r2, cn=CN_FONT, size=size, bold=True)
    r3 = p.add_run(text)
    set_run_font(r3, size=size)
    return p


def shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcpr.append(shd)


def cell_text(cell, text, size=9, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)


def add_image(doc, filename, caption, url, page, max_w=5.4, max_h=3.15):
    path = os.path.join(IMG, filename)
    if not os.path.exists(path):
        para(doc, "[缺失图片：%s]" % filename, size=9, color=GRAY)
        return
    with Image.open(path) as im:
        w, h = im.size
    width = max_w
    if h / w * width > max_h:
        width = max_h * w / h
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(width))

    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(1)
    cp.paragraph_format.keep_with_next = True
    r = cp.add_run("图：%s" % caption)
    set_run_font(r, cn=CN_FONT, size=9, bold=True, color=RGBColor(0x33, 0x33, 0x33))

    lp = doc.add_paragraph()
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lp.paragraph_format.space_after = Pt(10)
    r1 = lp.add_run("原图链接：")
    set_run_font(r1, size=8, color=GRAY)
    add_hyperlink(lp, url, url, size=8)
    r2 = lp.add_run("　|　来源页面：")
    set_run_font(r2, size=8, color=GRAY)
    add_hyperlink(lp, page, page, size=8)


# ---------------------------------------------------------------- 文档设置
def setup(doc):
    st = doc.styles["Normal"]
    st.font.name = EN_FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_BODY)
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.top_margin = Inches(0.9)
    sec.bottom_margin = Inches(0.9)

    # 页脚：文字 + 真实页码域
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    r = fp.add_run("广东工业大学机器人专业介绍 · 资料汇编　—　第 ")
    set_run_font(r, size=8, color=GRAY)

    def add_field(paragraph, instr):
        run = paragraph.add_run()
        set_run_font(run, size=8, color=GRAY)
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr_el = OxmlElement("w:instrText")
        instr_el.set(qn("xml:space"), "preserve")
        instr_el.text = " %s " % instr
        fld_sep = OxmlElement("w:fldChar")
        fld_sep.set(qn("w:fldCharType"), "separate")
        t = OxmlElement("w:t")
        t.text = "1"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        for el in (fld_begin, instr_el, fld_sep, t, fld_end):
            run._r.append(el)

    add_field(fp, "PAGE")
    r4 = fp.add_run(" 页　—　共 ")
    set_run_font(r4, size=8, color=GRAY)
    add_field(fp, "NUMPAGES")
    r5 = fp.add_run(" 页")
    set_run_font(r5, size=8, color=GRAY)


# ---------------------------------------------------------------- 论文表解析
def parse_papers():
    rows = []
    dist = []
    path = os.path.join(BASE, "论文佐证.md")
    with open(path, encoding="utf-8") as f:
        lines = f.read().splitlines()
    in_dist = False
    for ln in lines:
        s = ln.strip()
        if s.startswith("| ") and ("https://doi.org" in s or "robot.sia.cn" in s):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if len(cells) < 6:
                continue
            no, title, authors, venue, year, link = cells[0], cells[1], cells[2], cells[3], cells[4], cells[5]
            m = re.search(r"\((https?://[^)]+)\)", link)
            url = m.group(1) if m else ""
            rows.append(dict(no=no, title=title, authors=authors, venue=venue,
                             year=year, url=url, note=cells[6] if len(cells) > 6 else ""))
        if s.startswith("**方向分布速览**"):
            in_dist = True
            continue
        if in_dist:
            if s.startswith("- "):
                dist.append(s[2:].strip())
            elif s.startswith("---"):
                in_dist = False
    return rows, dist


# ---------------------------------------------------------------- 正文
def build():
    papers, dist = parse_papers()
    doc = Document()
    setup(doc)

    # ---------------- 封面
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("广东工业大学")
    set_run_font(r, cn=CN_FONT, size=30, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run("机 器 人 专 业 介 绍")
    set_run_font(r, cn=CN_FONT, size=26, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    r = p.add_run("Robotics Engineering, Guangdong University of Technology")
    set_run_font(r, size=12, color=GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    r = p.add_run("文 字 汇 编 　·　 图 片 资 料 　·　 论 文 佐 证")
    set_run_font(r, cn=CN_FONT, size=13, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run("资料索引见同目录 readme.md ／ 全部图片与文字均附原始来源链接")
    set_run_font(r, size=10, color=GRAY)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("广东工业大学　自动化学院 · 粤港机器人联合学院（机器人工程专业）")
    set_run_font(r, cn=CN_FONT, size=10.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("资料整理日期：2026 年 10 月")
    set_run_font(r, size=10, color=GRAY)

    doc.add_page_break()

    # ---------------- 编制说明
    h1(doc, "编制说明")
    para(doc, "本文档为「广东工业大学机器人专业」的资料汇编，面向专业认知、招生宣传、报考参考与教研备课等用途。"
              "全部文字均摘编自学校官网、学院官网、官方媒体报道等公开渠道，并在每一章节末尾标注原始链接；"
              "全部图片均下载自原始发布页面，标注原图地址与来源页面；论文部分逐条给出 DOI 解析链接。")
    para(doc, "体例约定：正文采用「官方原文摘编 + 来源标注」的方式，凡加引号的表述均为原文用语；"
              "凡数据（人数、比例、奖项、排名、经费）均照录原页面口径，未作换算。"
              "如不同页面口径存在差异（例如平台获批年份），文中同时标注两处来源，以学校层面明细表为准。")
    para(doc, "资料包结构：")
    bullet(doc, "广东工业大学机器人专业介绍.docx（本文档）——文字汇编 + 配图 + 论文佐证", bold_head="① ")
    bullet(doc, "readme.md——资料索引（逐文件、逐图、逐段来源说明）", bold_head="② ")
    bullet(doc, "论文佐证.md——28 篇代表性论文与科研平台的详细佐证材料", bold_head="③ ")
    bullet(doc, "images/——20 张配图（jpg/png），文件名即索引序号", bold_head="④ ")
    source_line(doc, "来源：广东工业大学官网 ", "https://www.gdut.edu.cn/")

    # ---------------- 第一章
    h1(doc, "第一章　专业概况")
    h2(doc, "1.1　专业定位：从「自动执行」到「自主智能」")
    para(doc, "机器人工程是广东工业大学面向国家战略性新兴产业重点建设的交叉复合型专业。"
              "学校官方专业介绍将机器人工程概括为「具身智能的物理容器」与「智慧产线的肌肉与骨骼」——"
              "它是让机器从「自动执行」迈向「自主智能」的核心专业，深度交叉机械、电子、控制、计算机与人工智能，"
              "属于典型的「新工科」。")
    para(doc, "与传统的机械或自动化专业相比，机器人工程的特点是「整机思维」："
              "学生既要懂机械本体怎么设计，也要懂电机怎么驱动、传感器怎么读、算法怎么让机器人「看懂」环境并「决定」下一步动作，"
              "最终把机械、电子、控制、软件整合成一台能完成真实任务的机器人。")
    source_line(doc, "来源：广东工业大学自动化学院《机器人工程专业介绍》（2021-12-27）",
                "https://automation.gdut.edu.cn/info/1035/8983.htm")

    h2(doc, "1.2　发展历程：从粤港机器人联合学院到机器人工程专业")
    para(doc, "广东工业大学的机器人人才培养起步于 2015 年的跨学科办学试验，2020 年正式获批本科专业，"
              "形成了「学院（荣誉学院式跨学科平台）＋ 专业（教育部备案本科专业）」双轨并行的格局。")
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0].cells
    for i, t in enumerate(["时间", "标志性事件", "意义"]):
        cell_text(hdr[i], t, size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(hdr[i], "DCE6F1")
    timeline = [
        ("2015 年 7 月", "联合香港科技大学李泽湘教授团队、东莞松山湖国际机器人产业基地，成立「粤港机器人联合学院」",
         "机器人工程专业的实体前身；首创「2+2」跨学科、项目驱动培养模式"),
        ("2017 年", "控制科学与工程在教育部第四轮学科评估中获 A-（全国前 10%）",
         "成为华南地区信息类唯一 A 类学科，为专业提供学科支撑"),
        ("2020 年", "「机器人工程」本科专业正式获教育部批准设立",
         "从「学院试验班」走向「国家专业目录」的规范化办学"),
        ("2021 年", "机器人工程入选广东省一流本科专业建设点",
         "省级层面获得专业建设认可"),
        ("2022 年", "第五轮学科评估续写优异成绩",
         "学科优势持续巩固"),
        ("2025 年", "软科中国最好学科排名：控制科学与工程位列全国第 20 名",
         "学科整体实力进入全国前列"),
    ]
    for t, e, m in timeline:
        cells = tbl.add_row().cells
        cell_text(cells[0], t, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(cells[1], e, size=9)
        cell_text(cells[2], m, size=9)
    source_line(doc, "来源：自动化学院《机器人工程专业介绍》 ", "https://automation.gdut.edu.cn/info/1035/8983.htm")
    source_line(doc, "　　　自动化学院《学科情况》 ", "https://automation.gdut.edu.cn/xkysz/xkjs.htm")
    source_line(doc, "　　　创新创业学院《以粤港澳机器人联合学院为代表的欧林工学院模式》 ",
                "https://cxcy.gdut.edu.cn/info/1148/2351.htm")

    h2(doc, "1.3　专业名片")
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    card = [
        ("专业名称", "机器人工程（Robotics Engineering）"),
        ("专业代码", "080803T（教育部本科专业目录）"),
        ("所属门类", "工学 · 自动化类"),
        ("办学主体", "自动化学院（专业）；粤港机器人联合学院（跨学科培养平台，2015 年成立）"),
        ("学制学位", "四年制本科，授予工学学士学位"),
        ("培养模式", "「2+2」：前两年在原专业学院打基础，后两年进入粤港机器人联合学院跨专业混编培养；全程项目驱动（PBL）"),
        ("依托学科", "控制科学与工程（教育部第四轮 A-、广东省攀峰重点学科、一级学科博士点、博士后流动站）"),
        ("交叉支撑", "机械工程、计算机科学与技术、信息与通信工程、材料科学与工程、生物医学工程、认知科学、工业设计、数学"),
        ("招生方式", "按「自动化类」大类招生，入学后第一学期在自动化类内选拔进入机器人工程专业"),
        ("专业荣誉", "2021 年入选广东省一流本科专业建设点；依托国家级科技企业孵化器、国家级众创空间等双创平台"),
    ]
    for k, v in card:
        cells = tbl.add_row().cells
        cell_text(cells[0], k, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(cells[0], "F2F2F2")
        cell_text(cells[1], v, size=9)
    source_line(doc, "来源：自动化学院《机器人工程专业介绍》 ", "https://automation.gdut.edu.cn/info/1035/8983.htm")
    source_line(doc, "　　　招生办《机器人工程（2022）》 ", "https://zsb.gdut.edu.cn/info/2334/2986.htm")
    source_line(doc, "　　　创新创业学院《粤港机器人联合学院简介》 ", "https://cxcy.gdut.edu.cn/info/1087/1428.htm")

    h2(doc, "1.4　学科支撑：控制科学与工程 A-")
    para(doc, "机器人工程专业的学科母体是控制科学与工程。该学科是广东省攀峰重点学科，"
              "拥有一级学科博士学位授权点、电子信息专业学位博士授权点及博士后流动站；"
              "在教育部第四轮学科评估中获评 A-（全国前 10%），第五轮评估「办学实力再获提升」；"
              "2025 年软科中国最好学科排名位列全国第 20 名，是「全国省属高校中该学科唯一的 A 类学科」，"
              "也是华南地区控制领域的核心优势学科。")
    para(doc, "学院另拥有控制科学与工程、电气工程 2 个一级学科硕士点，控制工程、电气工程、人工智能 3 个专业学位硕士领域；"
              "「电力系统及其自动化」为广东省特色重点学科。")

    para(doc, "学科发展脉络清晰：1996 年开始招收硕士研究生，1999 年招收「控制理论与控制工程」博士研究生，"
              "2003 年设立博士后流动站；2017 年教育部第四轮学科评估获 A-，「2020 软科」排名进入全国前 5%。"
              "学科现有博士生导师 35 名、硕士生导师 99 名；2015 年以来引进「百人计划」特聘教授 17 人、"
              "「青年百人」48 人（全部具有海外学习或工作经历），国家高层次 A 类人才 5 人，签约海外兼职专家 31 名。"
              "2019 年全球「高被引科学家」评选中，广东工业大学 6 人次入选，居内地高校第 22 位。")
    para(doc, "学科团队与平台支撑同样扎实：拥有教育部创新团队 1 个、科技部重点领域创新团队 1 个、"
              "广东省自然科学基金创新团队 3 个、全国高校「黄大年式」教师团队 1 个；"
              "并建有教育部重点实验室、国家「111 计划」学科创新引智基地等国家级平台（详见第三章）。")
    para(doc, "研究生课程体系中已开设与机器人直接相关的课程，例如「机器人控制技术」"
              "（课程编号 204138，32 学时／2 学分），以及运动控制系统、机器视觉识别与检测技术、"
              "模式识别与人工智能、智能控制基础、线性系统理论等，构成本科「机器人工程」专业向上衔接的课程通道。")
    source_line(doc, "来源：自动化学院《学科情况》 ", "https://automation.gdut.edu.cn/xkysz/xkjs.htm")
    source_line(doc, "　　　自动化学院《学术型硕士-控制科学与工程》（2025-09-01） ",
                "https://automation.gdut.edu.cn/info/1044/8867.htm")

    doc.add_page_break()

    # ---------------- 第二章
    h1(doc, "第二章　培养方案")
    h2(doc, "2.1　培养目标")
    para(doc, "专业培养目标共 4 条，指向「能解决复杂工程问题的机器人领域高素质创新型人才」：")
    goals = [
        ("家国情怀与职业素养：", "具有家国情怀、良好的人文社会科学素养、社会责任感与工程职业道德，能够在工程实践中理解并遵守工程规范。"),
        ("系统思维与复杂工程问题：", "能够综合运用数学、自然科学、工程基础与机器人专业知识，识别、表达并解决机器人及相关领域的复杂工程问题。"),
        ("沟通协调与团队领导：", "具备多学科团队中的沟通、协作与组织管理能力，能够在跨学科项目中承担负责人角色。"),
        ("自主学习与终身学习：", "具有自主学习和终身学习意识，能够跟踪机器人技术前沿并持续更新知识结构。"),
    ]
    for head, body in goals:
        bullet(doc, body, bold_head=head)
    source_line(doc, "来源：自动化学院《机器人工程专业介绍》·培养目标 ",
                "https://automation.gdut.edu.cn/info/1035/8983.htm")

    h2(doc, "2.2　毕业要求（12 条）")
    para(doc, "专业按工程教育专业认证的 12 项通用标准设置毕业要求，覆盖知识、能力、素养三个维度：")
    reqs = [
        "工程知识：能够将数学、自然科学、工程基础和专业知识用于解决复杂机器人工程问题。",
        "问题分析：能够应用数学、自然科学和工程科学的基本原理，识别、表达并通过文献研究分析复杂工程问题。",
        "设计/开发解决方案：能够设计针对复杂工程问题的解决方案，设计满足特定需求的机器人系统、部件或工艺流程。",
        "研究：能够基于科学原理并采用科学方法对复杂工程问题进行研究，包括设计实验、分析与解释数据。",
        "使用现代工具：能够针对复杂工程问题，开发、选择与使用恰当的技术、资源、现代工程工具和信息技术工具。",
        "工程与社会：能够基于工程相关背景知识进行合理分析，评价专业工程实践和复杂工程问题解决方案对社会、健康、安全、法律及文化的影响。",
        "环境和可持续发展：能够理解和评价针对复杂工程问题的工程实践对环境、社会可持续发展的影响。",
        "职业规范：具有人文社会科学素养、社会责任感，能够在工程实践中理解并遵守工程职业道德和规范。",
        "个人和团队：能够在多学科背景下的团队中承担个体、团队成员以及负责人的角色。",
        "沟通：能够就复杂工程问题与业界同行及社会公众进行有效沟通和交流，包括撰写报告和设计文稿、陈述发言、清晰表达。",
        "项目管理：理解并掌握工程管理原理与经济决策方法，并能在多学科环境中应用。",
        "终身学习：具有自主学习和终身学习的意识，有不断学习和适应发展的能力。",
    ]
    for i, t in enumerate(reqs, 1):
        bullet(doc, t, bold_head="%d. " % i)
    source_line(doc, "来源：自动化学院《机器人工程专业介绍》·毕业要求 ",
                "https://automation.gdut.edu.cn/info/1035/8983.htm")

    h2(doc, "2.3　课程体系：「三大基础、一个领域」")
    para(doc, "粤港机器人联合学院的课程体系被概括为「三大基础、一个领域」：人文社科基础、自然科学基础、专业基础合计 118 学分（占 73.8%），"
              "专业领域 42 学分（占 26.2%）。学院强调「数学基础、科学基础、工程基础与通识教育」，突出各专业核心课程与多专业融合项目。")
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0].cells
    for i, t in enumerate(["课程模块", "学分", "课时", "占比"]):
        cell_text(hdr[i], t, size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(hdr[i], "DCE6F1")
    credit = [
        ("公共基础", "59", "944", "36.88%"),
        ("专业基础", "28", "448", "17.50%"),
        ("专业课", "15", "240", "9.38%"),
        ("实验实习实训", "13", "208", "8.13%"),
        ("设计（论文）", "33", "528", "20.63%"),
        ("全校公选", "12", "200", "7.50%"),
        ("合计", "≥160", "≥2568", "100%"),
    ]
    for row in credit:
        cells = tbl.add_row().cells
        for i, v in enumerate(row):
            cell_text(cells[i], v, size=9,
                      align=WD_ALIGN_PARAGRAPH.CENTER if i else WD_ALIGN_PARAGRAPH.LEFT,
                      bold=(row[0] == "合计"))
    source_line(doc, "来源：创新创业学院《以粤港澳机器人联合学院为代表的欧林工学院模式》 ",
                "https://cxcy.gdut.edu.cn/info/1148/2351.htm")

    para(doc, "")
    h3(doc, "综合训练课程（贯穿四年）")
    for t in ["机器人入门项目设计Ⅰ／Ⅱ", "嵌入式系统设计", "机器人驱动与运动控制",
              "机器人机电一体项目设计", "机器人开放式项目设计", "机器人创新创业项目设计"]:
        bullet(doc, t)
    source_line(doc, "来源：创新创业学院《以粤港澳机器人联合学院为代表的欧林工学院模式》 ",
                "https://cxcy.gdut.edu.cn/info/1148/2351.htm")

    h3(doc, "课程设置的四条主线（以「机器人工程专业-专业课程」图示为准）")
    para(doc, "在 2025 年 3 月广州大学来访交流的专题报告中，粤港机器人联合学院责任教授黄之峰用一张分支图概括了机器人工程专业的课程骨架："
              "以「自动控制原理」为根，向上分出「传感」「运动控制」「计算机」「算法」四个方向——")
    for head, body in [
        ("传感：", "信号与系统、电路、模拟及数字电子技术、图像处理与机器视觉。"),
        ("运动控制：", "电路驱动与运动控制、Adams 建模与仿真实验。"),
        ("计算机：", "机器人操作系统、程序设计、计算机网络。"),
        ("算法：", "人工智能原理、算法与数据结构、机器学习。"),
    ]:
        bullet(doc, body, bold_head=head)
    para(doc, "该图同时标注「机器人操作系统／机器人学／机器人技术」为专业核心，落款为「核心专业课」。")
    source_line(doc, "来源：创新创业学院《广州大学到访我校，共探粤港机器人联合学院创新发展》（2025-03-13） ",
                "https://cxcy.gdut.edu.cn/info/1112/3522.htm")

    h2(doc, "2.4　「2+2」培养模式与全程项目驱动")
    para(doc, "粤港机器人联合学院采用「2+2」产教结合、校企融合的培养模式：前两年在学校，侧重数学基础、科学基础和专业基础的学习与工程能力培养，"
              "新生入学即引入「机器人导论」项目实践课程；后两年在东莞松山湖国际机器人产业基地培养，"
              "依托香港科技大学的师资力量、产业界高端人才与基地产业化能力，与机器人企业深度融合，"
              "侧重各专业与机器人紧密联系的领域课程学习、跨专业实际开发项目训练，并接受面向产业的金融课程、创业课程学习和训练，完成毕业设计。")
    para(doc, "项目驱动是这套模式的核心机制，具体安排如下：")
    for head, body in [
        ("大一 · 机器人入门项目设计Ⅰ（80 学时）：", "学生自己组队、自主动手制作一款最为基础的机器人，「真正喜欢上机器人」。"),
        ("大二 · 机器人入门项目设计Ⅱ（128 学时）：", "以团队形式围绕 RoboCon 竞赛主题展开，是入门项目Ⅰ的进阶版。"),
        ("大三／大四 · 3 个项目设计 + 毕业设计：", "每学期 1 个项目设计，面向复杂工程问题、紧密融合于机器人产业链各环节。"),
        ("假期特训营：", "在东莞松山湖国际机器人产业基地等地定期开设，为期一个半月的高强度集中训练，"
                        "通过集中授课和项目分组两种模式开展，并配合全国大学生机器人大赛 ROBOCON 实战练兵。"),
    ]:
        bullet(doc, body, bold_head=head)
    para(doc, "学院的教学理念被概括为「探究型教学，以项目为引导，以能力为导向」：以机器人及智能装备为对象，设计和引入实际工程项目，"
              "组织多专业项目团队，倡导「做中学、做中悟」，通过项目完成各专业知识点学习和能力培养；"
              "并秉承以学生为中心、成果导向教育（OBE）和持续质量改进（CQI）的教育理念。")
    source_line(doc, "来源：创新创业学院《粤港机器人联合学院简介》 ", "https://cxcy.gdut.edu.cn/info/1087/1428.htm")
    source_line(doc, "　　　中国教育报《广东工业大学机器人学院多专业融合、全程项目驱动——破壁生机育「新」才》（2024-10-07） ",
                "https://www.gdut.edu.cn/info/2045/21886.htm")

    h3(doc, "PBL 教学与交叉课程")
    para(doc, "机器人工程专业以智能制造为主要方向，在专业内再遴选学术创新、卓越工程师、机器人方面人才；"
              "安排工程管理、经济决策等交叉课程，通过工业 4.0 产线融合大数据、云计算、网络化，"
              "以自动化系列综合课程设计为载体，采用 PBL（Project/Problem-Based Learning）教学模式。"
              "专业学科基础涵盖控制科学与工程、机械工程、计算机科学与技术、材料科学与工程、生物医学工程、认知科学，"
              "研究领域包括智能感知、优化控制与系统设计、人机交互模式。")
    source_line(doc, "来源：招生办《机器人工程（2022）》 ", "https://zsb.gdut.edu.cn/info/2334/2986.htm")

    h3(doc, "「私人定制」的跨学科课程结构")
    para(doc, "学院 6 个学院的教师联合行业企业，综合原学科专业及跨学科专业知识，为机器人学院学生量身定制人才培养方案："
              "按原学科专业 65%、跨学科专业 35% 的比例为学生「私人定制」课程体系。"
              "每名学生都有自己的主体专业，而共同的专业方向都是机器人；"
              "跨学科专业课程以培养跨学科基础知识为主，强调学科和专业的跨界思维融合。")
    para(doc, "学院每年在大一新生中招收 60 至 80 名学生，学生来自 6 个实体学院的自动化等 7 个专业，按学科知识交叉、团队合作、动态调整的理念集中编班。"
              "（注：另一处官方表述为从机电工程、自动化、计算机、信息工程、艺术与设计、数学与统计 6 个学院招生，两说指同一批 6 个学院，专业计数口径略有差异。）")
    source_line(doc, "来源：中国教育报（2024-10-07） ", "https://www.gdut.edu.cn/info/2045/21886.htm")
    source_line(doc, "　　　创新创业学院《以粤港澳机器人联合学院为代表的欧林工学院模式》 ",
                "https://cxcy.gdut.edu.cn/info/1148/2351.htm")

    h2(doc, "2.5　学什么、做什么：一份能力地图")
    para(doc, "官方专业介绍把机器人工程的学习内容拆成四块，这套拆分同时也是一份「能力地图」：")
    for head, body in [
        ("机械本体：", "机械设计基础、机器人运动学与动力学；机械臂构型、足式底盘、轮足复合底盘的结构设计与选型。"),
        ("传感与驱动：", "电机与伺服驱动、液压与气动；激光雷达、深度相机、力矩传感器等感知器件的选型、标定与数据融合。"),
        ("控制与智能：", "自动控制原理、嵌入式系统、机器人操作系统（ROS）、SLAM、机器视觉、路径规划与运动控制。"),
        ("编程与数学：", "C++／Python 编程，数学建模与优化方法，算法实现与仿真验证。"),
    ]:
        bullet(doc, body, bold_head=head)
    para(doc, "毕业后的典型去向被归纳为四类：")
    for head, body in [
        ("工业机器人：", "机器人本体厂商、系统集成商，从事产线设计、部署、离线编程与调试。"),
        ("服务与特种机器人：", "扫地机器人、物流配送车、巡检无人机、应急救援装备等产品的研发与落地。"),
        ("医疗与人形机器人：", "手术机器人、康复外骨骼、具身智能与人形机器人整机研发。"),
        ("核心零部件：", "高精度伺服驱动、力／触觉传感器等关键部件的设计与制造。"),
    ]:
        bullet(doc, body, bold_head=head)
    para(doc, "官方建议：机器人工程适合「空间想象力强、动手能力突出、不怕多学科交叉」的学生；"
              "建议尽早选定主攻方向（运动控制／感知算法／结构设计），"
              "并借助 RoboMaster、RoboCup 等竞赛或实验室项目锻造实战能力。"
              "在职业发展上，既可以纵深成为运动控制算法专家，也可以横向成长为系统架构师。")
    source_line(doc, "来源：自动化学院《机器人工程专业介绍》 ", "https://automation.gdut.edu.cn/info/1035/8983.htm")

    doc.add_page_break()

    # ---------------- 第三章
    h1(doc, "第三章　师资队伍与科研平台")
    h2(doc, "3.1　师资规模")
    para(doc, "控制科学与工程学科现有博士生导师 35 名、硕士生导师 99 名；"
              "2015 年以来引进「百人计划」特聘教授 17 人、青年百人 48 人（全部具有海外履历）；"
              "国家高层次 A 类人才 5 人；2019 年全球「高被引科学家」评选中，广东工业大学 6 人次入选，居内地高校第 22 位；"
              "签约海外兼职专家 31 名。")
    para(doc, "团队建设方面，拥有教育部创新团队 1 个、科技部重点领域创新团队 1 个、广东省自然科学基金创新团队 3 个、"
              "全国高校「黄大年式」教师团队 1 个。")
    source_line(doc, "来源：自动化学院《学术型硕士-控制科学与工程》（2025-09-01） ",
                "https://automation.gdut.edu.cn/info/1044/8867.htm")

    h2(doc, "3.2　机器人方向主要学术带头人")
    leaders = [
        ("管贻生（Yisheng Guan）", "机电工程学院教授、博士生导师，仿生与智能机器人实验室（BIRL）创建人，"
                                   "广东工业大学智能智造与机器人研究所（系）所长，广东省「扬帆计划」高逼真仿人服务机器人创新团队负责人。"
                                   "研究方向：模块化仿生机器人、攀爬机器人、软体机器人、机器人作业与机器人抛光。"
                                   "据 OpenAlex 统计，累计发表论文 342 篇、被引 3562 次、h-index 28、i10-index 99。",
         "https://openalex.org/A5013471002"),
        ("苏春翌（Chun-Yi Su）", "自动化学院教授，国家特聘专家、教育部高层次人才讲座教授、「广东特支计划」杰出人才（南粤百杰）。"
                                 "方向为机器人及精密机电系统控制；发表论文 400 余篇，SCI 引用 8200 余次，Google Scholar 引用超 13800 次，h-index 59；"
                                 "曾任 IEEE Transactions on Automatic Control、IEEE Transactions on Control Systems Technology 副主编。",
         "https://www.gdut.edu.cn/info/1959/11590.htm"),
        ("鲁仁全（Renquan Lu）", "教授、博士生导师，广东工业大学党委常委、副校长；国家杰出青年科学基金获得者、「长江学者」特聘教授、"
                                 "万人计划领军人才、国家自然科学基金创新研究群体负责人、广东省重点实验室主任、全球高被引科学家。"
                                 "主要从事智能决策与协同控制及其在无人自主系统中的应用研究；发表 SCI 论文 150 余篇"
                                 "（含 Automatica 10 余篇、IEEE 汇刊 97 篇），SCI 他引 10000 余次。",
         "https://teacher.gdut.edu.cn/lurenquan/zh_CN/index.htm"),
        ("吴元清（Yuanqing Wu）", "教授、博士生导师，国家优秀青年科学基金获得者（2019）、国家重点研发计划项目首席科学家、"
                                  "广东省杰出青年基金获得者、广东省青年珠江学者。研究方向为人工智能算法、多机器人系统、四足机器人；"
                                  "在 IEEE TAC、Automatica 等发表论文 50 余篇，出版 2 部 Springer 英文专著。",
         "https://teacher.gdut.edu.cn/wuyuanqing/zh_CN/index/124395/list/"),
        ("任志刚（Zhigang Ren）", "副教授、博士生导师，知行具能实验室（CogniMotion Lab）主任，"
                                  "广东省离散制造知识自动化工程技术研究中心负责人，粤港澳离散制造智能化联合实验室骨干。"
                                  "研究方向为具身智能与机器人（无人车、四足／轮足机器人、移动操作机器人、机械臂抓取）；"
                                  "发表 SCI/EI 论文 100 余篇，ESI 高被引 3 篇、热点论文 2 篇，申请发明专利 70 余项（授权 30 项）。",
         "https://teacher.gdut.edu.cn/renzhigang/zh_CN/index.htm"),
        ("肖涵臻（Hanzhen Xiao）", "副教授、硕士生导师，广东工业大学「青年百人计划」A 层次引进人才，"
                                   "隶属智能检测与控制研究团队（自动化装备与集成重点实验室，章云教授领衔）。"
                                   "研究方向为智能移动机器人控制、多智能体编队、模型预测控制、视觉移动机器人、强化学习控制、无人机；"
                                   "发表 SCI/EI 论文 50 余篇；承担粤港机器人联合学院本科教学与 ROBOCON 大赛带队工作。",
         "https://teacher.gdut.edu.cn/xiaohanzhen/zh_CN/index.htm"),
    ]
    for name, desc, url in leaders:
        h3(doc, name)
        para(doc, desc)
        source_line(doc, "来源：", url)

    h2(doc, "3.3　国家级 / 部级科研平台")
    para(doc, "依托自动化学院建设的国家级、部级平台主要包括：")
    nat = [
        ("智能检测与制造物联教育部重点实验室", "2019 年，教育部，负责人谢胜利", "https://idiotm.gdut.edu.cn/"),
        ("物联网智能信息处理与系统集成国际合作联合实验室", "2015 年，教育部，负责人谢胜利",
         "https://www.gdut.edu.cn/kxyj/kyjgbf/sbj.htm"),
        ("国家学科创新引智基地（「111 计划」）——基于物联网技术的离散制造智能化学科创新引智基地", "2017 年，教育部、国家外国专家局",
         "https://automation.gdut.edu.cn/kxyj1/kypt.htm"),
        ("智能制造信息物理融合系统集成技术国家地方联合工程研究中心", "2016 年，国家发改委",
         "https://automation.gdut.edu.cn/kxyj1/kypt.htm"),
        ("粤港智能决策与协同控制联合实验室", "2023 年，教育部国际合作联合实验室（学校平台总表记作「粤港智能制造与协同控制联合实验室」）",
         "https://www.gdut.edu.cn/kxyj/kyjg1/ptljg1/zrkxlpt1/sbj.htm"),
        ("集成电路软硬件一体化教育部工程研究中心", "2022 年，教育部",
         "https://automation.gdut.edu.cn/kxyj1/kypt.htm"),
        ("低空感知探测技术与智能物联教育部工程研究中心", "2025 年，教育部（与低空经济、无人机感知方向直接相关）",
         "https://automation.gdut.edu.cn/kxyj1/kypt.htm"),
        ("国家智能制造试点示范基地：注塑成型智能装备与智能服务", "2015 年，工信部，中国制造 2025 专项",
         "https://automation.gdut.edu.cn/kxyj1/kypt.htm"),
    ]
    for name, meta, url in nat:
        bullet(doc, "（%s）" % meta, bold_head=name + "　")
        source_line(doc, "　　来源：", url)

    h2(doc, "3.4　省级科研平台")
    prov = [
        ("粤港澳离散制造智能化联合实验室", "2019 年，广东省科技厅，负责人谢胜利",
         "https://www.gdut.edu.cn/kxyj/kyjgbf/sbj.htm"),
        ("广东省物联网信息技术重点实验室", "2011 年，广东省科技厅，负责人谢胜利",
         "https://www.gdut.edu.cn/kxyj/kyjgbf/sbj.htm"),
        ("广东省信息物理融合系统重点实验室", "2015 年，广东省科技厅，负责人程良伦（学院平台页记 2016 年）",
         "https://www.gdut.edu.cn/kxyj/kyjgbf/sbj.htm"),
        ("广东省智能决策与协同控制重点实验室", "2018 年，广东省科技厅，负责人鲁仁全",
         "https://www.gdut.edu.cn/kxyj/kyjgbf/sbj.htm"),
        ("广东省高等学校重点实验室：自动化装备与集成实验室", "2001 年，广东省教育厅，章云教授领衔",
         "https://automation.gdut.edu.cn/kxyj1/kypt.htm"),
        ("广东省物联网共性技术研发工程中心", "2011 年，广东省经信委，负责人谢胜利",
         "https://automation.gdut.edu.cn/kxyj1/kypt.htm"),
        ("广东省车载传感网与智能交通工程技术研究中心", "2014 年，负责人谢胜利",
         "https://automation.gdut.edu.cn/kxyj1/kypt.htm"),
        ("广东省面向领域应用的大数据工程技术研究中心", "2015 年，负责人周郭许",
         "https://automation.gdut.edu.cn/kxyj1/kypt.htm"),
        ("广东省面向智能制造的超高精度自动化检测系统工程技术研究中心", "2016 年，负责人何昭水",
         "https://automation.gdut.edu.cn/kxyj1/kypt.htm"),
        ("广东省离散制造知识自动化工程技术研究中心", "2018 年，负责人吴宗泽",
         "https://automation.gdut.edu.cn/kxyj1/kypt.htm"),
    ]
    for name, meta, url in prov:
        bullet(doc, "（%s）" % meta, bold_head=name + "　")
    source_line(doc, "来源：广东工业大学《省部级科研平台明细表》 ", "https://www.gdut.edu.cn/kxyj/kyjgbf/sbj.htm")
    source_line(doc, "　　　自动化学院《科研平台》 ", "https://automation.gdut.edu.cn/kxyj1/kypt.htm")

    h2(doc, "3.5　代表性实验室")
    h3(doc, "仿生与智能机器人实验室（BIRL）")
    para(doc, "BIRL（Biomimetic and Intelligent Robotics Lab）隶属机电工程学院，由管贻生教授创建，"
              "是学校机器人方向最活跃的研究单元之一。研究方向分为五条线：仿生软体、模块攀爬、共融操作、SLAM、智能工业。"
              "实验室与广东富利盛仿生机器人股份有限公司等企业建立了战略合作，已研发出仿人机器人「艾比 AIBE」。"
              "BIRL 的实验室名称与「机电工程学院」归属直接出现在其论文的作者单位字段中（见第四章论文 1、4、6、7、11）。")
    source_line(doc, "来源：BIRL 实验室官网 ", "https://brilab.github.io/")
    source_line(doc, "　　　BIRL 研究方向 ", "https://brilab.github.io/research/")

    h3(doc, "智能检测与制造物联教育部重点实验室")
    para(doc, "2019 年由教育部批准建设，面向「卡脖子」技术攻关，设四大研究方向：盲信号智能检测、精密检测技术与装备、"
              "物联网智能信息处理理论与技术、制造物联技术与系统。2020 年发表高水平 SCI 论文 98 篇、中文核心论文 32 篇、授权发明专利 50 项。")
    source_line(doc, "来源：实验室官网 ", "https://idiotm.gdut.edu.cn/")

    h3(doc, "共融机器人视觉实验室（工大创谷双创团队）")
    para(doc, "团队拥有成员 30 余名，包括微软中国研究院首批研究员、广东省第五批引进创新创业团队带头人、"
              "国际模式识别学会会士刘文印教授，以及两名国际资深技术顾问和两名外籍博士后。"
              "实验室以提高机器人的「智能化」程度为目标，让机器人能够根据人体的动作变化进行安全的实时跟踪协作，"
              "并在环境发生巨大改变时通过视觉编程快速适应新工作环境，无需重新返厂编程。"
              "团队已申请 23 项发明专利，覆盖视频流学习、动作模仿、语音操作等领域。")
    source_line(doc, "来源：创新创业学院《共融机器人视觉实验室-创业团队》（2020-09-21） ",
                "https://cxcy.gdut.edu.cn/info/1137/1521.htm")

    h3(doc, "机器人技术实验室（信息工程学院实验中心）")
    para(doc, "实验室位于工学 1 号馆 403，面积 154 ㎡，主要面向通信工程、电子信息工程和信息工程等专业本科生与研究生的综合设计实验教学。"
              "实验室配备 AI 智能控制实训平台，能实现基于视觉、运动、语音等控制的机器人应用；"
              "主要设备包括 AI 视觉机械臂、机械臂控制套件、视频存储管理服务器、微型电子计算机等；"
              "开设电子信息系统综合设计 I-IV、信息系统综合设计 I-IV、自主驾驶与智能车竞赛、本科生创新创业训练等实验项目。")
    source_line(doc, "来源：信息工程学院实验中心《机器人技术实验室》 ",
                "https://xxxysyzx.gdut.edu.cn/bkjxsys/jqrjssys.htm")

    doc.add_page_break()

    # ---------------- 第四章
    h1(doc, "第四章　科研实力与论文佐证")
    h2(doc, "4.1　学科评估与科研获奖")
    para(doc, "控制科学与工程在教育部第四轮学科评估中获 A-，第五轮评估办学实力再获提升，2025 软科排名全国第 20 名。"
              "科研获奖方面，学院获国家自然科学二等奖 1 项、教育部一等奖 5 项／二等奖 2 项、"
              "广东省科学技术一等奖 4 项／二等奖 4 项；教学方面获国家级教学成果二等奖 1 项。")
    source_line(doc, "来源：自动化学院《学术型硕士-控制科学与工程》 ", "https://automation.gdut.edu.cn/info/1044/8867.htm")

    h2(doc, "4.2　代表性论文（28 篇，均附 DOI 链接）")
    para(doc, "下表收录 28 篇可核验的代表性论文，全部通过 Crossref 或 OpenAlex 逐条核对标题、作者、期刊与年份，"
              "作者单位字段可直接佐证广东工业大学相关实验室与平台的存在。"
              "详细版本（含作者全名单、单位标注、与专业的关联说明）见同目录《论文佐证.md》。")
    tbl = doc.add_table(rows=1, cols=5)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0].cells
    for i, t in enumerate(["#", "论文标题", "期刊 / 会议", "年份", "链接"]):
        cell_text(hdr[i], t, size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(hdr[i], "DCE6F1")
    for row in papers:
        cells = tbl.add_row().cells
        cell_text(cells[0], row["no"], size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(cells[1], row["title"], size=8)
        cell_text(cells[2], row["venue"], size=8)
        cell_text(cells[3], row["year"], size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        cells[4].text = ""
        pp = cells[4].paragraphs[0]
        pp.paragraph_format.space_after = Pt(1)
        pp.paragraph_format.space_before = Pt(1)
        short = row["url"].replace("https://doi.org/", "doi:").replace("https://", "")
        add_hyperlink(pp, row["url"], short, size=7.5)
    source_line(doc, "来源：论文元数据经 Crossref（https://api.crossref.org）与 OpenAlex（https://api.openalex.org）逐条核验",
                "https://openalex.org/I139024713")

    h2(doc, "4.3　方向分布")
    para(doc, "28 篇论文覆盖机器人学科的八个方向，可据此判断学校机器人研究的布局重心：")
    for d in dist:
        bullet(doc, d)
    source_line(doc, "来源：本资料包《论文佐证.md》整理 ", "")

    h2(doc, "4.4　从论文看学科特色")
    para(doc, "从上述论文可以看出广东工业大学机器人研究的三个特色：")
    for head, body in [
        ("仿生与模块化是传统强项：", "Climbot 双臂攀爬机器人、模块化双足爬壁机器人、轮足复合机器人、仿蚂蟥／仿尺蠖软体机器人，"
                                     "构成了一条从机构创新到步态规划的完整研究链，成果发表于 ASME JMR、IEEE/ASME TMech、ICRA、Journal of Bionic Engineering 等。"),
        ("控制理论底色厚：", "机械臂自适应控制、柔性关节预设性能控制、多智能体一致性、事件触发控制等成果密集发表于 IEEE TAC／TIE／TCYB／TNNLS／TSMC，"
                             "与「控制科学与工程 A-」的学科定位高度一致。"),
        ("与制造业深度绑定：", "机器人抛光力控末端、AGV 实时最优导航、输电线路巡检机器人、微纳机器人精密制造等，"
                               "都指向广东「制造业当家」的产业需求，也是「智能检测与制造物联」教育部重点实验室的主攻方向。"),
    ]:
        bullet(doc, body, bold_head=head)

    doc.add_page_break()

    # ---------------- 第五章
    h1(doc, "第五章　学生培养成效")
    h2(doc, "5.1　竞赛成绩")
    para(doc, "学院实施「项目驱动和竞赛驱动」双轮驱动模式，以赛促练促学。")
    h3(doc, "ROBOCON（全国大学生机器人大赛）")
    para(doc, "连续 5 年参加 ROBOCON，2020 年获全国一等奖（季军），与东北大学一起代表中国参加当年 ABU 亚太机器人大赛（ABU Festival）国际赛事。"
              "下图为 ROBOCON 2020 季军奖杯与全国一等奖证书。")
    h3(doc, "ROBOMASTER（机甲大师赛）")
    para(doc, "2020 年，DynamicX 战队在第十九届全国大学生机器人大赛 ROBOMASTER 2020 机甲大师对抗赛（线上）中获一等奖 2 项、三等奖 6 项："
              "对抗赛获三等奖；步兵机器人研发获全国第 8 名（一等奖）；嵌入式研发获全国第 12 名（一等奖）。"
              "该年共 308 所高校参赛，包括宾州州立大学、纽约大学、阿尔伯塔大学、新加坡国立大学、哈尔滨工业大学、华南理工大学、"
              "上海交通大学、电子科技大学等。指导老师为鲁仁全、吴元清；队员来自自动化、物理与光电、机电工程、计算机、信息工程、"
              "材料与能源、国际教育、外国语等学院。赛事由共青团中央、深圳市人民政府主办，DJI 大疆承办，2020 年列入广工一级竞赛名录。")
    source_line(doc, "来源：广工新闻《我校学子在2020年全国大学生机器人大赛ROBOMASTER2020中获一等奖》（2020-10-10） ",
                "https://www.gdut.edu.cn/info/1709/11727.htm")
    source_line(doc, "　　　创新创业学院《以粤港澳机器人联合学院为代表的欧林工学院模式》 ",
                "https://cxcy.gdut.edu.cn/info/1148/2351.htm")

    h3(doc, "其他赛事")
    para(doc, "在「挑战杯」、中国国际「互联网+」大学生创新创业大赛等赛事中，机器人学院斩获近 200 项国家级、省级奖项；"
              "另一处（2021 年口径）统计为「互联网+」「创青春」「挑战杯」等获省级、国家级以上奖项近 40 项。"
              "两处数据统计口径与时间不同，前者为 2024 年累计口径。")
    source_line(doc, "来源：中国教育报（2024-10-07） ", "https://www.gdut.edu.cn/info/2045/21886.htm")

    h2(doc, "5.2　专利与论文")
    para(doc, "学生申请专利 30 余件，发表高水平论文 10 篇（2024 年口径）；2021 年口径为申请专利 20 余项、公开发表学术论文近 10 篇。")
    source_line(doc, "来源：中国教育报（2024-10-07） ", "https://www.gdut.edu.cn/info/2045/21886.htm")

    h2(doc, "5.3　毕业生去向")
    para(doc, "截至 2024 年，机器人学院已培养 213 名毕业生，其中国内外升学读研占 40%、就业创业占 60%。"
              "不少毕业生在百度、字节跳动、大疆等知名企业就业。"
              "深造院校包括米兰理工大学、中国科学院大学、香港中文大学、西安交通大学、电子科技大学、华南理工大学等。")
    h3(doc, "毕业生案例")
    for head, body in [
        ("黄裕超（大疆研发工程师）：", "作为骨干成员之一参与仰望 U8 车载无人机项目的硬件开发，"
                                       "通过团队不断优化设计方案，成功将这款车载无人机推向市场。他评价：「学校跨学科的学习经历为我在工作中"
                                       "进行跨部门、跨领域合作，奠定了坚实的基础，使我在面对复杂问题时，能够组织多模块、不同领域的同事共同解决问题。」"),
        ("陈栩生（信息工程学院）：", "大一、大二在原专业学习电子信息类课程和机器人工程专业融合课程，大三、大四在机器人学院与不同专业同学共同学习。"
                                     "在无人机研发课程上负责设计供电与驱动模块——自动化专业学生负责控制算法，机械类成员承担结构设计，工业设计学生负责外观设计。"),
        ("庄栢杨：", "参与腰椎穿刺设备、智能肿瘤识别仪等医疗机器人研发，在医疗机器人临床一线开展工作。"),
        ("黄泽凯（2024 届机械方向）：", "通过特训营以及 ROBOCON 赛事实战练兵，提升了团队合作、动手、组织等能力；"
                                        "现就职于深圳拓竹科技有限公司，其负责设计的 3D 打印机零件已大规模量产。"),
    ]:
        bullet(doc, body, bold_head=head)
    source_line(doc, "来源：中国教育报《广东工业大学机器人学院多专业融合、全程项目驱动——破壁生机育「新」才》（2024-10-07） ",
                "https://www.gdut.edu.cn/info/2045/21886.htm")
    source_line(doc, "　　　创新创业学院《以粤港澳机器人联合学院为代表的欧林工学院模式》 ",
                "https://cxcy.gdut.edu.cn/info/1148/2351.htm")

    doc.add_page_break()

    # ---------------- 第六章
    h1(doc, "第六章　招生与报考")
    h2(doc, "6.1　招生方式")
    para(doc, "机器人工程按「自动化类」大类招生，入学后于第一学期在自动化类内选拔进入机器人工程专业。"
              "粤港机器人联合学院则每年在大一新生中招收 60–80 人，覆盖 6 个实体学院的 7 个专业，"
              "属于「二次选拔＋跨学院混编」的培养通道。")
    source_line(doc, "来源：教育部阳光高考平台在线咨询（2024） ",
                "https://gaokao.chsi.com.cn/zxdy/forum--method-listDefault,forumid-52327,year-2024,start-465.dhtml")
    source_line(doc, "　　　中国教育报（2024-10-07） ", "https://www.gdut.edu.cn/info/2045/21886.htm")

    h2(doc, "6.2　2025 年新增：机器人与机电一体化工程国际班")
    para(doc, "2025 年学校新增「机械设计制造及其自动化（机器人与机电一体化工程国际班）」，与澳大利亚蒙纳士大学学分互认，首届招生。"
              "该班与机器人工程专业形成互补：前者偏机械本体与国际联合培养，后者偏控制与智能、依托自动化学院。")
    source_line(doc, "来源：国际教育学院 ", "https://sie.gdut.edu.cn/info/1012/6226.htm")

    h2(doc, "6.3　什么样的学生适合读")
    para(doc, "综合官方表述，机器人工程适合以下特质的学生：")
    for t in ["空间想象力强，对机械结构与运动有直觉；",
              "动手能力突出，愿意拆装、调试、写代码；",
              "不怕多学科交叉，能同时接受机械、电子、控制、计算机的训练；",
              "有长期主义心态，愿意尽早选定一个主攻方向并持续深耕。"]:
        bullet(doc, t)
    para(doc, "官方给出的路径建议是：借 RoboMaster、RoboCup 等竞赛或实验室项目锻造实战能力，"
              "在本科阶段就把「整机做出来一次」，这比单纯刷绩点更能决定后续发展空间。")
    source_line(doc, "来源：自动化学院《机器人工程专业介绍》 ", "https://automation.gdut.edu.cn/info/1035/8983.htm")

    doc.add_page_break()

    # ---------------- 第七章 图片
    h1(doc, "第七章　图片资料（20 张，均附原图链接）")
    para(doc, "本章收录 20 张与机器人专业直接相关的图片，全部下载自学校官网、学院官网与实验室官网的原始发布页面。"
              "每张图片下方标注图说、原图地址与来源页面；图片文件同时保存在同目录 images/ 文件夹中，文件名即本章序号。")

    images = [
        ("01_xinhua_huawei_sign.jpg",
         "广东工业大学与华为共建鲲鹏、昇腾开发者创享社揭牌（新华社 2022-09-16 摄）",
         "https://www.gdut.edu.cn/__local/D/6D/B4/516480AFE87DE98DFEAB5620D15_28E87732_152FD.jpg",
         "https://www.gdut.edu.cn/info/2045/16630.htm"),
        ("02_xinhua_coding.jpg",
         "学生在「工大创谷」创新创业工作室编写程序（新华社 2022-09-16 摄）",
         "https://www.gdut.edu.cn/__local/D/C7/E7/33D67CC1A18D711A14F60F48FB9_977D012F_282BD.jpg",
         "https://www.gdut.edu.cn/info/2045/16630.htm"),
        ("03_xinhua_amphibious_platform.jpg",
         "学生团队在「工大创谷」研发空水两栖远程巡检平台（新华社 2022-09-16 摄）",
         "https://www.gdut.edu.cn/__local/B/18/81/757507A6661A8B5AE72DA9F656F_F538E496_1853E.jpg",
         "https://www.gdut.edu.cn/info/2045/16630.htm"),
        ("04_xinhua_robot_research.jpg",
         "广东工业大学粤港机器人联合学院的学生在「工大创谷」进行机器人研发（新华社 2022-09-16 摄）",
         "https://www.gdut.edu.cn/__local/9/0B/A7/85DDE878E884194FC1E7EC1A44E_597937C5_1A446.jpg",
         "https://www.gdut.edu.cn/info/2045/16630.htm"),
        ("05_xinhua_vision_coding.jpg",
         "「工大创谷」创客工坊团队的梁梓熙在编写计算机视觉深度学习程序（新华社 2022-09-16 摄）",
         "https://www.gdut.edu.cn/__local/0/11/24/715C1F83C34EA34F40C6D938B9B_9EB205C2_1C31D.jpg",
         "https://www.gdut.edu.cn/info/2045/16630.htm"),
        ("06_robomaster2020_a.jpg",
         "ROBOMASTER 2020 一等奖证书两张：步兵机器人研发（全国第 8）、嵌入式研发（全国第 12）",
         "https://www.gdut.edu.cn/__local/8/29/93/BA638FC489EF28A709081AA4338_7D964C51_11242.jpg",
         "https://www.gdut.edu.cn/info/1709/11727.htm"),
        ("07_robomaster2020_b.jpg",
         "ROBOMASTER 2020 机甲大师对抗赛（线上）三等奖证书",
         "https://www.gdut.edu.cn/__local/C/3F/8A/02DD8B8660C8902C57BA431B3A4_28450D14_1D876.jpg",
         "https://www.gdut.edu.cn/info/1709/11727.htm"),
        ("08_robotics_lab_visit_a.jpg",
         "粤港机器人联合学院责任教授黄之峰讲解「机器人工程专业-专业课程」分支图"
         "（自动控制原理为根，分出传感／运动控制／计算机／算法四支）",
         "https://cxcy.gdut.edu.cn/__local/0/50/52/2759A48825F83919B43DC4B1496_47B860FB_A5F6.jpg",
         "https://cxcy.gdut.edu.cn/info/1112/3522.htm"),
        ("09_robotics_lab_visit_b.jpg",
         "参观人员合影，背景为「粤港机器人联合学院 Robotics Union Academy Guangdong-Hong Kong」标识墙",
         "https://cxcy.gdut.edu.cn/__local/B/44/E1/8AFBB1959D88217BAB7B2359B34_D77FE4EE_10627.jpg",
         "https://cxcy.gdut.edu.cn/info/1112/3522.htm"),
        ("10_robotics_academy_sign.png",
         "粤港机器人联合学院门头标识墙（校徽 + 蓝盾 R 标志 + 中英文院名）",
         "https://cxcy.gdut.edu.cn/__local/6/11/F1/57C6B457BDE4A411743F7CE07AF_BBFE1408_92769.png",
         "https://cxcy.gdut.edu.cn/info/1148/2351.htm"),
        ("11_robocon_award.png",
         "ROBOCON 2020 季军奖杯与全国一等奖证书（第十九届全国大学生机器人大赛 ROBOCON 2020 绿茵争锋线上赛）",
         "https://cxcy.gdut.edu.cn/__local/2/22/E4/0AEAF10DA3621C170F0F8F0CE0D_759432A9_9F453.png",
         "https://cxcy.gdut.edu.cn/info/1148/2351.htm"),
        ("12_chinaedu_robot_college.jpg",
         "《中国教育报》2024 年 10 月 7 日版面：广东工业大学机器人学院多专业融合、全程项目驱动——破壁生机育「新」才",
         "https://www.gdut.edu.cn/__local/5/5D/D4/E7A875E5937E95D0D429ED4C395_31655675_10666C3.png",
         "https://www.gdut.edu.cn/info/2045/21886.htm"),
        ("13_robot_tech_lab.png",
         "信息工程学院「机器人技术实验室」AI 智能控制实训平台：①平台组成 ②平台技术结构体系 ③平台设备与开发套件 ④平台应用场景",
         "https://xxxysyzx.gdut.edu.cn/__local/4/C8/C4/4EF1CFF4F12234F84D2D6264944_2753A2A7_DF702.png",
         "https://xxxysyzx.gdut.edu.cn/bkjxsys/jqrjssys.htm"),
        ("14_collab_robot_pour.png",
         "共融机器人视觉实验室成果：机器人模仿人体动作倒水（视觉编程 + 实时跟踪协作）",
         "https://cxcy.gdut.edu.cn/__local/2/98/BF/C1E7D1FC5557229DA9624E95C28_65F1CACB_D79F.png",
         "https://cxcy.gdut.edu.cn/info/1137/1521.htm"),
        ("15_collab_robot_microwave.jpg",
         "共融机器人视觉实验室成果：机器人操作微波炉（面向家庭场景的共融操作）",
         "https://cxcy.gdut.edu.cn/__local/5/E0/3E/16339FCB98A1E496DFF1A50DCBA_FE683689_F3F5.png",
         "https://cxcy.gdut.edu.cn/info/1137/1521.htm"),
        ("16_birl_soft_robotics.jpg",
         "BIRL 研究方向之一「仿生软体」：模块化软体臂、柔性驱动器与多指软体手（含抓取生鸡蛋等脆弱物体演示）",
         "https://brilab.github.io/images/research/soft-robotics.jpg",
         "https://brilab.github.io/research/"),
        ("17_birl_module_climbing.jpg",
         "BIRL 研究方向「模块攀爬／模块化机器人」官网配图：机械臂抓取 3D 打印异形模块构件",
         "https://brilab.github.io/images/research/module-climbing.jpg",
         "https://brilab.github.io/research/"),
        ("18_birl_human_robotics.jpg",
         "BIRL 研究方向「共融操作」官网配图：Robot Introspection（机器人内省）技术路线——"
         "由力／位姿／声音等多模态信号经状态估计与异常监测，支撑决策、故障恢复与运动生成，实现鲁棒的机器人自主操作",
         "https://brilab.github.io/images/research/human-robotics.jpg",
         "https://brilab.github.io/research/"),
        ("19_birl_slam.jpg",
         "BIRL 研究方向「SLAM」官网配图：Segway RMP 移动机器人平台，搭载视觉传感器、笔记本与手柄",
         "https://brilab.github.io/images/research/slam.jpg",
         "https://brilab.github.io/research/"),
        ("20_birl_intelligent_industry.jpg",
         "BIRL 研究方向「智能工业」官网配图：五杆并联机构（并联机器人）构型与尺度综合示意",
         "https://brilab.github.io/images/research/intelligent-industry.jpg",
         "https://brilab.github.io/research/"),
    ]
    for fn, cap, url, page in images:
        h3(doc, "图 %s　%s" % (fn[:2], fn))
        add_image(doc, fn, cap, url, page)

    doc.add_page_break()

    # ---------------- 第八章 参考资料
    h1(doc, "第八章　参考资料清单")
    h2(doc, "8.1　学校与学院官方页面")
    refs = [
        ("自动化学院《机器人工程专业介绍》（2021-12-27）", "https://automation.gdut.edu.cn/info/1035/8983.htm"),
        ("自动化学院《学科情况》", "https://automation.gdut.edu.cn/xkysz/xkjs.htm"),
        ("自动化学院《学术型硕士-控制科学与工程》（2025-09-01）", "https://automation.gdut.edu.cn/info/1044/8867.htm"),
        ("自动化学院《科研平台》", "https://automation.gdut.edu.cn/kxyj1/kypt.htm"),
        ("招生办《机器人工程（2022）》", "https://zsb.gdut.edu.cn/info/2334/2986.htm"),
        ("创新创业学院《粤港机器人联合学院简介》", "https://cxcy.gdut.edu.cn/info/1087/1428.htm"),
        ("创新创业学院《以粤港澳机器人联合学院为代表的欧林工学院模式》",
         "https://cxcy.gdut.edu.cn/info/1148/2351.htm"),
        ("创新创业学院《广州大学到访我校，共探粤港机器人联合学院创新发展》（2025-03-13）",
         "https://cxcy.gdut.edu.cn/info/1112/3522.htm"),
        ("创新创业学院《共融机器人视觉实验室-创业团队》（2020-09-21）", "https://cxcy.gdut.edu.cn/info/1137/1521.htm"),
        ("信息工程学院实验中心《机器人技术实验室》", "https://xxxysyzx.gdut.edu.cn/bkjxsys/jqrjssys.htm"),
        ("广东工业大学《省部级科研平台明细表》", "https://www.gdut.edu.cn/kxyj/kyjgbf/sbj.htm"),
        ("广东工业大学《省部级科研平台分类总表》", "https://www.gdut.edu.cn/kxyj/kyjg1/ptljg1/zrkxlpt1/sbj.htm"),
        ("智能检测与制造物联教育部重点实验室", "https://idiotm.gdut.edu.cn/"),
        ("苏春翌教授简介", "https://www.gdut.edu.cn/info/1959/11590.htm"),
        ("国际教育学院（2025 机器人与机电一体化工程国际班）", "https://sie.gdut.edu.cn/info/1012/6226.htm"),
        ("仿生与智能机器人实验室 BIRL", "https://brilab.github.io/"),
        ("BIRL 研究方向", "https://brilab.github.io/research/"),
    ]
    for name, url in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-12)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("▪ " + name + "：")
        set_run_font(r, size=9)
        add_hyperlink(p, url, url, size=8.5)

    h2(doc, "8.2　新闻与媒体")
    news = [
        ("中国教育报《广东工业大学机器人学院多专业融合、全程项目驱动——破壁生机育「新」才》（2024-10-07）",
         "https://www.gdut.edu.cn/info/2045/21886.htm"),
        ("中国教育报电子版原版", "http://paper.jyb.cn/zgjyb/html/2024-10/07/content_144745_17864725.htm"),
        ("新华网图片报道《广东工大产教融合实战育人培养双创人才》（2022-09-16）",
         "https://www.gdut.edu.cn/info/2045/16630.htm"),
        ("新华网原文", "http://www.news.cn/photo/2022-09/16/c_1129009075_5.htm"),
        ("广工新闻《我校学子在2020年全国大学生机器人大赛ROBOMASTER2020中获一等奖》（2020-10-10）",
         "https://www.gdut.edu.cn/info/1709/11727.htm"),
    ]
    for name, url in news:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-12)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("▪ " + name + "：")
        set_run_font(r, size=9)
        add_hyperlink(p, url, url, size=8.5)

    h2(doc, "8.3　教师主页与学术档案")
    teachers = [
        ("鲁仁全教授主页", "https://teacher.gdut.edu.cn/lurenquan/zh_CN/index.htm"),
        ("吴元清教授主页", "https://teacher.gdut.edu.cn/wuyuanqing/zh_CN/index/124395/list/"),
        ("任志刚副教授主页", "https://teacher.gdut.edu.cn/renzhigang/zh_CN/index.htm"),
        ("肖涵臻副教授主页", "https://teacher.gdut.edu.cn/xiaohanzhen/zh_CN/index.htm"),
        ("章云教授主页", "https://teacher.gdut.edu.cn/zhangyun/zh_CN/index/23627/list/index.htm"),
        ("管贻生 OpenAlex 作者页", "https://openalex.org/A5013471002"),
        ("广东工业大学 OpenAlex 机构页", "https://openalex.org/I139024713"),
        ("教育部第四轮学科评估结果（0811 控制科学与工程）",
         "https://www.cdgdc.edu.cn/dslxkpgjggb/xkpm/gx/a0811_kzkxygc.htm"),
    ]
    for name, url in teachers:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-12)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("▪ " + name + "：")
        set_run_font(r, size=9)
        add_hyperlink(p, url, url, size=8.5)

    h2(doc, "8.4　论文 DOI 链接（28 条）")
    for row in papers:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-12)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run("[%s] " % row["no"])
        set_run_font(r, size=8.5, bold=True)
        r2 = p.add_run(row["title"][:80] + ("…" if len(row["title"]) > 80 else "") + "（%s, %s）" % (row["venue"], row["year"]))
        set_run_font(r2, size=8.5)
        add_hyperlink(p, row["url"], " " + row["url"], size=8)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("— 全文完 —")
    set_run_font(r, cn=CN_FONT, size=11, color=GRAY)

    doc.save(OUT)
    print("saved:", OUT, os.path.getsize(OUT), "bytes")
    print("papers:", len(papers), "dist:", len(dist))


if __name__ == "__main__":
    build()
